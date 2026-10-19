import os
import unittest
import tempfile
from docx import Document

from src.core.models import ReportProject, DateGroup, ActivityItem
from src.core.docx_generator import DocxGenerator
from src.core.image_processor import ImageProcessor

class TestDocxGenerator(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()
        self.output_docx = os.path.join(self.temp_dir, "test_output.docx")

        # Ambil sampel foto asli jika ada
        sample_folder = r"F:\TK Cahaya Hati\Foto Timestamp Reevitalisasi TK Cahaya Hati 10-12-26\2026-10-12"
        self.sample_photos = []
        if os.path.exists(sample_folder):
            files = [os.path.join(sample_folder, f) for f in os.listdir(sample_folder) if f.lower().endswith(('.jpg', '.png'))]
            self.sample_photos = files[:4]

    def test_generate_report_structure(self):
        project = ReportProject(
            project_name="Test Laporan Revitalisasi",
            photo_width_cm=7.5
        )

        dg1 = DateGroup(date_str="2026-10-12")
        act1 = ActivityItem(
            description="Foto Bangunan Tampak Depan",
            photo_paths=self.sample_photos[:2]
        )
        act2 = ActivityItem(
            description="Foto Bangunan Tampak Samping Kanan",
            photo_paths=self.sample_photos[2:3]
        )
        dg1.activities.extend([act1, act2])

        dg2 = DateGroup(date_str="2026-10-13")
        act3 = ActivityItem(
            description="Foto Material yang Dibeli :\nSemen\nBatu Bata\nPasir",
            photo_paths=self.sample_photos[1:3]
        )
        dg2.activities.append(act3)

        project.date_groups.extend([dg1, dg2])

        # Generate Word
        out_file = DocxGenerator.generate_report(project, self.output_docx)
        self.assertTrue(os.path.exists(out_file))

        # Verifikasi isi docx
        doc = Document(out_file)
        self.assertEqual(len(doc.tables), 1)
        tbl = doc.tables[0]

        # Header + 2 activities (dg1) + 1 activity (dg2) = 4 rows
        self.assertEqual(len(tbl.rows), 4)

        # Header check
        self.assertEqual(tbl.rows[0].cells[0].text.strip(), "No.")
        self.assertEqual(tbl.rows[0].cells[1].text.strip(), "Tanggal")
        self.assertEqual(tbl.rows[0].cells[2].text.strip(), "Keterangan")
        self.assertEqual(tbl.rows[0].cells[3].text.strip(), "Dokumentasi")

        # Row 1 (dg1 first activity)
        self.assertEqual(tbl.rows[1].cells[0].text.strip(), "1.")
        self.assertTrue("Senin, 12 Oktober 2026" in tbl.rows[1].cells[1].text)
        self.assertEqual(tbl.rows[1].cells[2].text.strip(), "Foto Bangunan Tampak Depan")

        # Row 2 (dg1 second activity on same date)
        self.assertEqual(tbl.rows[2].cells[0].text.strip(), "")
        self.assertEqual(tbl.rows[2].cells[1].text.strip(), "")
        self.assertEqual(tbl.rows[2].cells[2].text.strip(), "Foto Bangunan Tampak Samping Kanan")

        # Row 3 (dg2 first activity on next date)
        self.assertEqual(tbl.rows[3].cells[0].text.strip(), "2.")
        self.assertTrue("Selasa, 13 Oktober 2026" in tbl.rows[3].cells[1].text)
        self.assertTrue("Foto Material yang Dibeli :" in tbl.rows[3].cells[2].text)

        print("[TEST SUCCESS] DocxGenerator test passed perfectly!")

if __name__ == "__main__":
    unittest.main()
