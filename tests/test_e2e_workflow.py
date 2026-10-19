import os
import unittest
import tempfile
import zipfile
import xml.etree.ElementTree as ET
from docx import Document

from src.core.models import ReportProject, DateGroup, ActivityItem
from src.core.docx_generator import DocxGenerator
from src.core.image_processor import ImageProcessor

class TestE2EWorkflow(unittest.TestCase):
    def test_full_workflow_and_persistence(self):
        sample_base = r"F:\TK Cahaya Hati\Foto Timestamp Reevitalisasi TK Cahaya Hati 10-12-26"
        self.assertTrue(os.path.exists(sample_base), "Sample base folder must exist")

        temp_dir = tempfile.mkdtemp()
        draft_json = os.path.join(temp_dir, "draft_project.json")
        output_docx = os.path.join(temp_dir, "Laporan_Revitalisasi_Final.docx")

        # 1. Bangun proyek
        project = ReportProject(
            project_name="Laporan Revitalisasi TK Cahaya Hati Oktober 2026",
            base_folder=sample_base,
            photo_width_cm=7.5
        )

        dates = ["2026-10-12", "2026-10-13", "2026-10-14"]
        for d in dates:
            dg = project.get_or_create_date_group(d)
            folder_d = os.path.join(sample_base, d)
            all_photos = [
                os.path.join(folder_d, f) for f in os.listdir(folder_d)
                if f.lower().endswith(('.jpg', '.png'))
            ]
            self.assertTrue(len(all_photos) > 5)

            if d == "2026-10-12":
                dg.activities.append(ActivityItem(
                    description="Foto Bangunan Tampak Depan",
                    photo_paths=all_photos[:2]
                ))
                dg.activities.append(ActivityItem(
                    description="Foto Bangunan Tampak Samping Kanan",
                    photo_paths=all_photos[2:4]
                ))
            elif d == "2026-10-13":
                dg.activities.append(ActivityItem(
                    description="Foto Material yang Dibeli :\nSemen\nBatu Bata\nBesi Cor\nPasir",
                    photo_paths=all_photos[:4]
                ))
            elif d == "2026-10-14":
                dg.activities.append(ActivityItem(
                    description="Foto Bangunan Tampak Belakang",
                    photo_paths=all_photos[:2]
                ))

        # 2. Test Simpan dan Muat Draft JSON
        project.save_to_json(draft_json)
        self.assertTrue(os.path.exists(draft_json))

        loaded_project = ReportProject.load_from_json(draft_json)
        self.assertEqual(len(loaded_project.date_groups), 3)
        self.assertEqual(len(loaded_project.date_groups[0].activities), 2)
        self.assertEqual(len(loaded_project.date_groups[1].activities), 1)

        # 3. Test Generate Dokumen Word
        generated_file = DocxGenerator.generate_report(loaded_project, output_docx)
        self.assertTrue(os.path.exists(generated_file))
        self.assertTrue(os.path.getsize(generated_file) > 100000) # File berisi gambar > 100KB

        # 4. Verifikasi Struktur Tabel Dokumen Word
        doc = Document(generated_file)
        self.assertEqual(len(doc.tables), 1)
        tbl = doc.tables[0]
        # Total rows: 1 Header + 2 (dg1) + 1 (dg2) + 1 (dg3) = 5 rows
        self.assertEqual(len(tbl.rows), 5)

        # Cek XML untuk skala gambar
        with zipfile.ZipFile(generated_file) as z:
            xml_content = z.read("word/document.xml")
            tree = ET.fromstring(xml_content)
            ns = {
                'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
                'wp': 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'
            }
            extents = tree.findall('.//wp:extent', ns)
            # Total gambar = 2 + 2 + 4 + 2 = 10 gambar
            self.assertEqual(len(extents), 10)
            for ext in extents:
                cx_cm = int(ext.get('cx', 0)) / 360000
                self.assertAlmostEqual(cx_cm, 7.5, delta=0.05)

        print(f"[E2E TEST PASSED] Word generated with 10 photos, exact table layout, size={os.path.getsize(generated_file)/1024:.1f} KB")

if __name__ == "__main__":
    unittest.main()
